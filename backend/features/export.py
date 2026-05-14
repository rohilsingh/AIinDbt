"""Export helpers — Excel via openpyxl, PDF via browser print."""
from __future__ import annotations

import io
from typing import Any

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False


def to_excel(rows: list[dict[str, Any]], columns: list[str], sheet_name: str = "Data") -> bytes:
    if not HAS_XLSX:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]

    header_fill = PatternFill("solid", fgColor="1E40AF")
    header_font = Font(bold=True, color="FFFFFF")

    for ci, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=ci, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(columns, 1):
            val = row.get(col, "")
            if isinstance(val, (list, dict)):
                val = str(val)
            ws.cell(row=ri, column=ci, value=val)
        if ri % 2 == 0:
            for ci in range(1, len(columns) + 1):
                ws.cell(row=ri, column=ci).fill = PatternFill("solid", fgColor="F1F5F9")

    for ci, col in enumerate(columns, 1):
        max_len = max(
            (len(str(row.get(col, ""))) for row in rows[:200]),
            default=0,
        )
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = min(
            max(len(col), max_len) + 4, 50
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
